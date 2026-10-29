"""Smoke tests for the leads Excel export endpoint and the offline
transcript-analyzer enrichment stored on leads."""
import io

import pytest
from openpyxl import load_workbook

from app.models.lead import Lead


@pytest.mark.asyncio
async def test_export_leads_excel(client, db_session):
    lead = Lead(
        phone_number="+919822012345",
        name="Rohit Sharma",
        email="rohit.sharma@gmail.com",
        status="interested",
        score=0.97,
        reason="Asked about price and agreed to a demo.",
        metadata_json={
            "customer_info": {"name": "Rohit Sharma", "city": "Pune",
                              "budget_inr": 200000, "product_interest": None},
            "why_interested": [{"code": "price_inquiry",
                                "reason": "Asked about price / cost",
                                "evidence": "What is the price?"}],
            "why_not_interested": [],
        },
    )
    other = Lead(phone_number="+919000000000", status="not_interested",
                 score=0.1, reason="Said not interested.", metadata_json={})
    db_session.add_all([lead, other])
    await db_session.commit()

    # All leads
    resp = await client.get("/api/leads/export.xlsx")
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Name"
    assert len(rows) == 3  # header + 2 leads
    flat = str(rows)
    assert "Rohit Sharma" in flat
    assert "Asked about price / cost" in flat
    assert "Pune" in flat

    # Filtered: interested only
    resp2 = await client.get("/api/leads/export.xlsx?status=interested")
    assert resp2.status_code == 200
    wb2 = load_workbook(io.BytesIO(resp2.content))
    rows2 = list(wb2.active.iter_rows(values_only=True))
    assert len(rows2) == 2  # header + 1 lead
    assert "not_interested" not in str(rows2[1])
